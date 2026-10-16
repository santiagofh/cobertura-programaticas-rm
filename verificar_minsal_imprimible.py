"""Verifica que el Excel imprimible MINSAL sea copia fiel del corte oficial.

Compara celda a celda las 3 hojas de
`output/coberturas_vacunas_{T}_MINSAL_imprimible.xlsx` (Cobertura %, Vacunados,
Población objetivo) contra los CSV oficiales del corte MINSAL, reconstruyendo
la matriz con las mismas reglas del exportador (etiquetas cortas, 52 comunas
ordenadas + fila TOTAL REGIÓN METROPOLITANA).

Además revisa:
  - El CSV semilla `coberturas_vacunas_{T}_V7_MINSAL.csv` contra el detalle
    por comuna del corte.
  - La suma comunal contra la fila nacional publicada por MINSAL (cuadre).

Uso:
    python verificar_minsal_imprimible.py
    python verificar_minsal_imprimible.py --corte-dir "ruta" --anio 2025 --tolerancia 1e-6

Salida: PASS/FAIL por bloque con el detalle de diferencias; código de salida 1
si algo no cuadra.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
import pandas as pd

from convertir_minsal_programaticas import CORTE_DEFAULT, cargar_corte
from exportar_excels_imprimibles import FILA_TOTAL, _preparar_pivots
from src import config_programaticas as cfgmod

HOJAS = [
    ("Cobertura %", "pct"),
    ("Vacunados", "vac"),
    ("Población objetivo", "den"),
]

fallas: list[str] = []


def _registrar(ok: bool, mensaje: str) -> None:
    print(f"  [{'OK' if ok else 'FALLA'}] {mensaje}")
    if not ok:
        fallas.append(mensaje)


def _comparar_celdas(hoja: str, df_excel: pd.DataFrame, esperado: pd.DataFrame,
                     tolerancia: float) -> None:
    diffs = []
    for comuna in esperado.index:
        for col in esperado.columns:
            vx = df_excel.at[comuna, col]
            ve = esperado.at[comuna, col]
            if pd.isna(ve) and (pd.isna(vx) or vx == "S/D"):
                continue
            if pd.isna(ve) != pd.isna(vx):
                diffs.append((comuna, col, vx, ve))
            elif abs(float(vx) - float(ve)) > tolerancia:
                diffs.append((comuna, col, vx, ve))
    _registrar(not diffs, f"{hoja}: {len(esperado.index) * len(esperado.columns)} "
                          f"celdas comparadas, {len(diffs)} diferencias")
    for d in diffs[:15]:
        print(f"      comuna={d[0]!r} columna={d[1]!r} excel={d[2]!r} esperado={d[3]!r}")
    if len(diffs) > 15:
        print(f"      ... y {len(diffs) - 15} diferencias más")


def verificar_imprimible(ruta: Path, det: pd.DataFrame, tolerancia: float) -> None:
    pct, vac, den, _, _ = _preparar_pivots(det)
    esperados = {"pct": pct, "vac": vac, "den": den}

    wb = openpyxl.load_workbook(ruta, data_only=True)
    _registrar(wb.sheetnames == [h for h, _ in HOJAS],
               f"hojas del Excel: {wb.sheetnames}")

    for hoja, clave in HOJAS:
        if hoja not in wb.sheetnames:
            _registrar(False, f"falta la hoja {hoja!r}")
            continue
        ws = wb[hoja]
        esperado = esperados[clave]
        cols = list(esperado.columns)
        cols_excel = [c.value for c in ws[2]][1:1 + len(cols)]
        _registrar(cols_excel == cols, f"{hoja}: encabezados de vacunas")
        if cols_excel != cols:
            print(f"      excel={cols_excel}")
            print(f"      esperado={cols}")

        filas = {}
        for row in ws.iter_rows(min_row=4, max_row=4 + len(esperado.index) - 1,
                                values_only=True):
            filas[str(row[0]).strip()] = list(row[1:1 + len(cols)])
        df_excel = pd.DataFrame.from_dict(filas, orient="index", columns=cols)
        _registrar(list(df_excel.index) == [str(i) for i in esperado.index],
                   f"{hoja}: 52 comunas + {FILA_TOTAL}")
        if list(df_excel.index) == [str(i) for i in esperado.index]:
            df_excel.index = esperado.index
            _comparar_celdas(hoja, df_excel, esperado, tolerancia)


def verificar_semilla(ruta_semilla: Path, det: pd.DataFrame, tolerancia: float) -> None:
    semilla = pd.read_csv(ruta_semilla, encoding="utf-8-sig")
    m = det.merge(semilla, on=["COD_COMUNA_JOIN", "VACUNA_DASHBOARD"],
                  suffixes=("_of", "_sem"), how="outer", indicator=True)
    _registrar((m["_merge"] == "both").all(),
               f"semilla {ruta_semilla.name}: {len(m)} filas cruzadas "
               f"({(m['_merge'] != 'both').sum()} sin par)")
    for col in ["VACUNAS_ADMINISTRADAS", "POBLACION_OBJETIVO"]:
        dif = (m[col + "_of"].astype(float) - m[col + "_sem"].astype(float)).abs()
        _registrar(dif.max() <= tolerancia,
                   f"semilla: {col} (max dif {dif.max():.0f})")
    dif_pct = (m["PORC_COBERTURA_of"] - m["PORC_COBERTURA_sem"]).abs()
    _registrar(dif_pct.max() <= tolerancia,
               f"semilla: PORC_COBERTURA (max dif {dif_pct.max():.4f})")


def verificar_cuadre_nacional(det: pd.DataFrame, nac: pd.DataFrame,
                              tolerancia: float) -> None:
    agg = det.groupby("VACUNA_DASHBOARD", as_index=False).agg(
        vacunados=("VACUNAS_ADMINISTRADAS", "sum"),
        poblacion=("POBLACION_OBJETIVO", "sum"),
    )
    m = agg.merge(nac[["vacuna_usuario", "minsal_vacunados", "minsal_poblacion"]],
                  left_on="VACUNA_DASHBOARD", right_on="vacuna_usuario", how="left")
    sin_fila = m["minsal_vacunados"].isna()
    _registrar((~sin_fila).all(),
               f"cuadre nacional: {sin_fila.sum()} vacunas sin fila nacional")
    dif_vac = (m["vacunados"] - m["minsal_vacunados"].fillna(m["vacunados"])).abs()
    dif_pob = (m["poblacion"] - m["minsal_poblacion"].fillna(m["poblacion"])).abs()
    _registrar(dif_vac.max() <= tolerancia and dif_pob.max() <= tolerancia,
               f"cuadre nacional: suma comunal vs fila publicada "
               f"(max dif vacunados={dif_vac.max():.0f}, población={dif_pob.max():.0f})")


def main():
    ap = argparse.ArgumentParser(
        description="Verifica que el imprimible MINSAL calce con la fuente oficial")
    ap.add_argument("--corte-dir", type=Path, default=CORTE_DEFAULT)
    ap.add_argument("--anio", type=int, default=2025)
    ap.add_argument("--imprimible", type=Path, default=None)
    ap.add_argument("--semilla", type=Path, default=None)
    ap.add_argument("--tolerancia", type=float, default=1e-6)
    args = ap.parse_args()

    imprimible = args.imprimible or (
        cfgmod.OUTPUT_DIR / f"coberturas_vacunas_{args.anio}_MINSAL_imprimible.xlsx")
    semilla = args.semilla or (
        cfgmod.OUTPUT_DIR / f"coberturas_vacunas_{args.anio}_V7_MINSAL.csv")

    print(f"Corte oficial : {args.corte_dir}")
    print(f"Imprimible    : {imprimible}")
    print(f"Semilla       : {semilla}")
    print()

    det, nac = cargar_corte(args.corte_dir)
    print(f"Detalle oficial: {len(det)} filas | "
          f"{det['VACUNA_DASHBOARD'].nunique()} combinaciones | "
          f"{det['COD_COMUNA_JOIN'].nunique()} comunas")
    print()

    print("1) Imprimible vs corte oficial")
    verificar_imprimible(imprimible, det, args.tolerancia)
    print()
    print("2) CSV semilla vs corte oficial")
    verificar_semilla(semilla, det, args.tolerancia)
    print()
    print("3) Cuadre suma comunal vs fila nacional")
    verificar_cuadre_nacional(det, nac, args.tolerancia)
    print()

    if fallas:
        print(f"RESULTADO: FALLA ({len(fallas)} chequeos con diferencias)")
        sys.exit(1)
    print("RESULTADO: PASS — el imprimible es copia fiel de la fuente oficial")


if __name__ == "__main__":
    main()
